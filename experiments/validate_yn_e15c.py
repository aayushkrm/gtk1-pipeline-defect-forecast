"""YN E15 phase 4: depth-sanity shares, offset/orient fill, dedup check,
pipe-number cross-check 2020-weld vs 2023-pipelog vs 2020-anom."""
import json
import csv
import collections
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "outputs"
BASE = REPO.parent / "Данные для предварительного изучения" / "Отчет ВТД_Ю-Н_0-154"


def num(x):
    try:
        return float(str(x).replace(",", "."))
    except (ValueError, TypeError):
        return None


def load(path):
    with open(path, encoding="utf-8") as f:
        r = csv.DictReader(f)
        return r.fieldnames, list(r)


def main():
    out = {}
    fn20, r20 = load(OUT / "yn2020_anom_salvaged.csv")
    fn3, f3 = load(OUT / "yn2023_feat_salvaged.csv")
    fn4, f4 = load(OUT / "yn2023_anom_salvaged.csv")

    # 1. depth sanity: share of measured d,% within (0,100]
    for tag, rows, c in [("feat", f3, "d,\n%"), ("anom", f4, "d,\n%")]:
        vals = [num(r[c]) for r in rows]
        ok = [v for v in vals if v is not None]
        sane = [v for v in ok if 0 < v <= 100]
        import statistics
        print(f"{tag} d%: measured={len(ok)} sane(0,100]={len(sane)} "
              f"({len(sane)/len(ok):.3f}) med_all={statistics.median(ok)} "
              f"med_sane={statistics.median(sane) if sane else None}")
        hist = collections.Counter(
            ("<=100" if v <= 100 else "<=1000" if v <= 1000 else
             "<=10000" if v <= 10000 else ">10000") for v in ok)
        print(f"  magnitude hist: {dict(hist)}")
        out[tag + "_ddepth_sane"] = {"measured": len(ok), "sane": len(sane)}
        # sane-depth rows: pipe/distance filled?
        srows = [r for r in rows if (lambda v: v is not None and 0 < v <= 100)(num(r[c]))]
        pf = sum(1 for r in srows if num(r["N\nтрубы"]) not in (None, 60462.0)
                 and num(r["N\nтрубы"]) < 20000)
        df = sum(1 for r in srows if num(r["Дистанция по\nодометру, м"]) not in (None, 60462.0)
                 and num(r["Дистанция по\nодометру, м"]) < 160000)
        print(f"  sane-depth rows={len(srows)} sane-pipe={pf} sane-dist={df}")
        out[tag + "_sane_rows"] = {"n": len(srows), "pipe": pf, "dist": df}

    # 2. offset / orient fill
    for tag, rows in [("feat", f3), ("anom", f4)]:
        off = [r["Расстояние от\nпоперечных швов, м"] for r in rows]
        on = [num(v) for v in off]
        ok = [v for v in on if v is not None]
        print(f"{tag} offset: filled={len(ok)}/{len(rows)} "
              f"min={min(ok) if ok else None} max={max(ok) if ok else None} "
              f"nonnum_sample={collections.Counter(v for v in off if num(v) is None).most_common(4)}")
        ori = [r["Ориент.,\nч:мин"] for r in rows]
        print(f"{tag} orient: filled={sum(1 for v in ori if v)}/{len(rows)}")
    for c in ["От левого шва до точки максимума, м",
              "Минимальное расстояние до продольного шва, мм",
              "От репера, м", "До репера, м"]:
        vals = [num(r[c]) for r in r20]
        ok = [v for v in vals if v is not None]
        print(f"2020 {c[:34]}: filled={len(ok)}/{len(vals)}")

    # 3. feat dedup: are repeated-r rows identical?
    seen = {}
    ident = diff = 0
    for r in f3:
        k = r["_row"]
        sig = json.dumps([r[c] for c in fn3[2:]], ensure_ascii=False)
        if k in seen:
            if seen[k] == sig:
                ident += 1
            else:
                diff += 1
        else:
            seen[k] = sig
    print(f"feat dup rows: identical={ident} differing={diff}")
    out["feat_dups"] = {"identical": ident, "differing": diff}

    # 4. pipe-number cross-check via weld logs
    import openpyxl
    wb = openpyxl.load_workbook(
        str(BASE / "2020" / "Поперечные сварные швы.xlsx"),
        read_only=True, data_only=True)
    ws = wb["Поперечные сварные швы"]
    it = iter(ws.iter_rows(values_only=True))
    w20 = [r for r in it]
    # find header row (SSID/Номер трубы/Расстояние)
    hidx = next(i for i, r in enumerate(w20)
                if r[1] == "Номер трубы" or (r[0] == "SSID" and r[1] == "Номер трубы"))
    wmap = {}
    for r in w20[hidx + 1:]:
        try:
            wmap[int(float(str(r[1])))] = float(str(r[2]).replace(",", "."))
        except (ValueError, TypeError, IndexError):
            pass
    print("2020 weld pipes:", len(wmap), "span:",
          min(wmap.values()), max(wmap.values()))
    wb3 = openpyxl.load_workbook(
        str(BASE / "2023" / "otchet_скорр.xlsx"),
        read_only=True, data_only=True)
    ws3 = wb3["Трубный журнал"]
    it3 = iter(ws3.iter_rows(values_only=True))
    next(it3)
    pmap = {}
    for r in it3:
        try:
            pmap[int(float(str(r[0])))] = float(str(r[1]).replace(",", "."))
        except (ValueError, TypeError):
            pass
    print("2023 pipelog pipes:", len(pmap), "span:",
          min(pmap.values()), max(pmap.values()))
    # same pipe numbers at same distances?
    common = sorted(set(wmap) & set(pmap))
    print("common pipe numbers:", len(common))
    diffs = [abs(wmap[k] - pmap[k]) for k in common[::max(1, len(common) // 2000)][:2000]]
    import statistics
    print(f"weld20-vs-pipe23 distance delta: med={statistics.median(diffs)} "
          f"max={max(diffs)} n={len(diffs)}")
    out["pipe_numbering"] = {"common": len(common),
                             "median_delta": statistics.median(diffs),
                             "max_delta": max(diffs)}
    # 2020 anomalies: pipe P at dist D vs weld20 pipe P
    hit = tot = 0
    deltas = []
    for r in r20:
        p, d = num(r["Номер трубы"]), num(r["Расстояние, м"])
        if p is None or d is None or int(p) not in wmap:
            continue
        tot += 1
        dd = abs(wmap[int(p)] - d)
        deltas.append(dd)
        if dd < 12:
            hit += 1
    deltas.sort()
    print(f"2020anom pipe-vs-weld20: n={tot} within12m={hit} "
          f"med_delta={deltas[len(deltas)//2] if deltas else None}")
    out["anom20_vs_weld20"] = {"n": tot, "within12m": hit,
                               "med": deltas[len(deltas)//2] if deltas else None}
    # 2020 anomalies vs 2023 pipelog numbering
    hit3 = tot3 = 0
    for r in r20:
        p, d = num(r["Номер трубы"]), num(r["Расстояние, м"])
        if p is None or d is None or int(p) not in pmap:
            continue
        tot3 += 1
        if abs(pmap[int(p)] - d) < 12:
            hit3 += 1
    print(f"2020anom pipe-vs-pipe23: n={tot3} within12m={hit3}")
    out["anom20_vs_pipe23"] = {"n": tot3, "within12m": hit3}
    # 2023 anom journal vs 2023 pipelog numbering
    hit4 = tot4 = 0
    for r in f4:
        p, d = num(r["N\nтрубы"]), num(r["Дистанция по\nодометру, м"])
        if p is None or d is None or int(p) not in pmap or p == 60462:
            continue
        tot4 += 1
        if abs(pmap[int(p)] - d) < 12:
            hit4 += 1
    print(f"2023anom pipe-vs-pipe23: n={tot4} within12m={hit4}")
    out["anom23_vs_pipe23"] = {"n": tot4, "within12m": hit4}

    (OUT / "yn_e15_phase4.json").write_text(json.dumps(out, indent=2, ensure_ascii=False))
    print("wrote yn_e15_phase4.json")


if __name__ == "__main__":
    main()
