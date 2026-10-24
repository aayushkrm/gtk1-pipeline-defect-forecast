#!/usr/bin/env python3
"""ingest_npk.py: read the Датасет НПК archive into one tidy table + audit.

Input:  /Users/akm/Tsu Project/Датасет НПК/Датасет НПК/<year>/<section>/Аномалии.xls*
Output: outputs/npk/anomalies.csv          (tidy rows, not committed)
        experiments/results_npk_audit.json (aggregates only, committed)

Schema mapping (fixed for this object):
  dist       <- "Расстояние, м"
  pipe_no    <- "Номер трубы"
  pipe_len   <- "Длина трубы, м"
  off_mm     <- "От продольного шва до центра, мм"
  h          <- "Ориентация максимума, ч:мин" (fallback "Ориентация центра, ч:мин")
  depth_pct  <- "Глубина, %"
  length_mm  <- "Длина, мм"
  width_mm   <- "Ширина, мм"
  wall_mm    <- "Толщина, мм"
  char       <- "Характер особенности"
  kind       <- "Тип особенности"
  size_class <- "Класс размера"
  desc       <- "Описание"
  danger     <- "Опасность"
  kb         <- "КБД"
  no_years   <- "Срок НО, лет"
  pf         <- "Pf, МПа"

Section identity comes from the "Общая информация" sheet: pipeline name, km range,
diameter, and the tool table (launch dates) -> survey_date.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path

import pandas as pd

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))
from gtk1.io import _apply_tolerant_xlrd  # noqa: E402

COLMAP = {
    "SSID": "ssid",
    "Расстояние": "dist",
    "Номер трубы": "pipe_no",
    "Длина трубы": "pipe_len",
    "От продольного шва до центра": "off_mm",
    "От левого шва до начала": "off",
    "Ориентация максимума": "ori_max",
    "Ориентация центра": "h",
    "Глубина": "depth_pct",
    "Длина": "length_mm",
    "Ширина": "width_mm",
    "Толщина": "wall_mm",
    "Характер особенности": "char",
    "Тип особенности": "kind",
    "Класс размера": "size_class",
    "Описание": "desc",
    "Опасность": "danger",
    "КБД": "kb",
    "Срок НО": "no_years",
    "Pf": "pf",
    "Pd": "pd",
    "MAOP": "maop",
    "Расположение": "location",
}

NUMERIC = ("dist", "pipe_no", "pipe_len", "off_mm", "h", "depth_pct",
           "length_mm", "width_mm", "wall_mm", "no_years", "pf", "pd", "maop")


def read_rows(fp: Path, sheet: str):
    """Return list-of-lists for one sheet; handles .xls content under any name."""
    if open(fp, "rb").read(2) == b"PK":
        import openpyxl

        with open(fp, "rb") as fh:
            wb = openpyxl.load_workbook(fh, read_only=True, data_only=True)
            if sheet not in wb.sheetnames:
                wb.close()
                return None
            ws = wb[sheet]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            wb.close()
            return rows

    import xlrd

    _apply_tolerant_xlrd()
    bk = xlrd.open_workbook(fp, on_demand=False)
    if sheet not in bk.sheet_names():
        return None
    sh = bk.sheet_by_name(sheet)
    return [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]


def header_index(rows, first_col: str) -> int | None:
    for i, row in enumerate(rows[:12]):
        for cell in row:
            if str(cell).strip().startswith(first_col):
                return i
    return None


def to_hours(v) -> float:
    """'ч:мин' orientation -> decimal hours; empty -> -1."""
    s = str(v).strip()
    if not s or s.lower() in ("nan", "none"):
        return -1.0
    m = re.match(r"^(\d{1,2}):(\d{2})", s)
    if not m:
        try:
            return float(s)
        except ValueError:
            return -1.0
    hh, mm = int(m.group(1)), int(m.group(2))
    if hh > 12:  # tolerate 13:00 style encodings as 01:00
        hh -= 12
    return hh + mm / 60.0


def num(v):
    s = str(v).strip().replace(",", ".").replace("\xa0", "").replace(" ", "")
    if not s or s.lower() in ("nan", "none"):
        return float("nan")
    try:
        return float(s)
    except ValueError:
        return float("nan")


def parse_section(folder: str):
    """'Основная труба (406-572 км)' -> ('main', 406.0, 572.0)."""
    kind = "main" if folder.lower().startswith("основная") else "branch"
    m = re.search(r"\(([^)]*)\)", folder)
    rng = (float("nan"), float("nan"))
    if m:
        txt = m.group(1).replace("км", "").replace(",", ".")
        nums = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", txt)]
        if len(nums) >= 2:
            rng = (min(nums[0], nums[1]), max(nums[0], nums[1]))
        elif len(nums) == 1:
            rng = (0.0, nums[0])
    return kind, rng[0], rng[1]


def parse_info(rows):
    """Key/value pairs from the Общая информация sheet + tool launch dates."""
    kv, tools = {}, []
    for row in rows or []:
        cells = [("" if c is None else str(c).strip()) for c in row]
        is_date = len(cells) >= 2 and bool(re.match(r"^\d{2}\.\d{2}\.\d{4}$", cells[1]))
        if cells and cells[0].lower().startswith("дата создания"):
            is_date = False
        if is_date:
            tools.append({"tool": cells[0], "launch": cells[1],
                          "accept": cells[2] if len(cells) > 2 else ""})
            continue
        if len(cells) >= 2 and cells[0] and cells[1] and cells[0] not in kv:
            kv[cells[0]] = cells[1]
    return kv, tools


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default="/Users/akm/Tsu Project/Датасет НПК/Датасет НПК")
    ap.add_argument("--out", default="outputs/npk/anomalies.csv")
    ap.add_argument("--audit", default="experiments/results_npk_audit.json")
    args = ap.parse_args()

    root = Path(args.root)
    frames, audit = [], []
    for fp in sorted(root.rglob("Аномалии*.xls*")):
        rel = fp.relative_to(root)
        year = int(rel.parts[0])
        folder = unicodedata.normalize("NFC", rel.parts[1])
        kind, km0, km1 = parse_section(folder)

        info_rows = read_rows(fp, "Общая информация")
        kv, tools = parse_info(info_rows)
        rows = read_rows(fp, "Аномалии")
        if rows is None:
            audit.append({"file": str(rel), "error": "no Аномалии sheet"})
            continue
        hi = header_index(rows, "Расстояние")
        if hi is None:
            audit.append({"file": str(rel), "error": "no header row"})
            continue
        header = [str(c).strip() for c in rows[hi]]
        cols = {}
        # Longest key first; one source column maps to at most one tidy name.
        # This keeps "Длина" from also claiming "Длина трубы".
        keys = sorted(COLMAP.items(), key=lambda kv: len(kv[0]), reverse=True)
        for j, h in enumerate(header):
            for key, name in keys:
                if h.startswith(key):
                    if name not in cols:
                        cols[name] = j
                    break
        body = rows[hi + 1:]
        recs = []
        for row in body:
            if not row or not str(row[0]).strip():
                continue
            rec = {}
            for name, j in cols.items():
                rec[name] = row[j] if j < len(row) else None
            recs.append(rec)
        df = pd.DataFrame(recs)
        if df.empty:
            audit.append({"file": str(rel), "rows": 0})
            continue
        for c in NUMERIC + ("off",):
            if c in df.columns:
                df[c] = df[c].map(num)
            else:
                df[c] = float("nan")
        for c in ("char", "kind", "size_class", "desc", "danger", "kb", "location"):
            df[c] = df.get(c, "").astype(str).str.strip() if c in df.columns else ""
        if "ssid" in df.columns:
            def _sid(v):
                s = str(v).strip()
                if s.endswith(".0"):
                    s = s[:-2]
                return s
            df["ssid"] = df["ssid"].map(_sid)
        else:
            df["ssid"] = ""
        df["h"] = df["h"].map(to_hours) if "h" in df.columns else -1.0
        if "ori_max" in df.columns:
            hmax = df["ori_max"].map(to_hours)
            df["h"] = hmax.where(hmax >= 0, df["h"])

        df["year"] = year
        df["section"] = folder
        df["kind_sec"] = kind
        df["km0"] = km0
        df["km1"] = km1
        df["pipeline"] = kv.get("Трубопровод", "")
        df["uchastok"] = kv.get("Участок", "")
        df["diameter"] = kv.get("Диаметр", "")
        df["survey_date"] = tools[0]["launch"] if tools else ""
        df["file"] = str(rel)
        frames.append(df)

        d = df["dist"].dropna()
        depth = df["depth_pct"].dropna()
        dup = int(df.duplicated(subset=["dist"]).sum())
        audit.append({
            "file": str(rel), "year": year, "section": folder, "kind": kind,
            "km_range": [km0, km1], "pipeline": df["pipeline"].iloc[0],
            "uchastok": df["uchastok"].iloc[0], "diameter": df["diameter"].iloc[0],
            "survey_date": df["survey_date"].iloc[0],
            "rows": int(len(df)),
            "dist": [float(d.min()), float(d.max())] if len(d) else None,
            "depth": {"min": float(depth.min()), "max": float(depth.max()),
                      "n_gt_20": int((depth > 20).sum()),
                      "n_gt_100": int((depth > 100).sum()),
                      "n_nan": int(df["depth_pct"].isna().sum())},
            "pipe_no": {"n_unique": int(df["pipe_no"].nunique()),
                        "min": float(df["pipe_no"].min()), "max": float(df["pipe_no"].max())},
            "dup_dist": dup,
            "chars": df["char"].value_counts().head(6).to_dict(),
            "danger": df["danger"].value_counts().head(5).to_dict(),
        })

    all_df = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    out = REPO / args.out
    out.parent.mkdir(parents=True, exist_ok=True)
    all_df.to_csv(out, index=False)
    (REPO / args.audit).write_text(json.dumps(audit, ensure_ascii=False, indent=1), encoding="utf-8")

    print(f"{len(all_df)} rows from {len(frames)} files -> {out}")
    for a in audit:
        if "rows" in a:
            print(f"{a['year']} {a['section'][:42]:42s} rows={a['rows']:6d} "
                  f"dist={a['dist'][0]:.0f}-{a['dist'][1]:.0f} "
                  f"depth_max={a['depth']['max']:.1f} >100%={a['depth']['n_gt_100']} "
                  f"pipes={a['pipe_no']['n_unique']} dup={a['dup_dist']} date={a['survey_date']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
